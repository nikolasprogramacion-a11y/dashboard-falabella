"""Lee el archivo de presupuesto (formato PRESUPUESTO_FINAL_X_MES_X_Q.xlsx).

Estructura esperada en la hoja (la primera del archivo):
  fila 1: etiquetas de trimestre (Q1, Q2...), combinadas en varias columnas
  fila con 'Periodo (calendario)' en la 1ra columna: mes de cada columna (ej. sep-26)
  fila con 'Ingresos' en la 1ra columna: el monto de presupuesto de cada mes
"""
import io
import openpyxl
import pandas as pd
import streamlit as st

_MESES = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12,
}


def periodo_a_fechas(periodo):
    """'sep-26' -> (2026-09-01, 2026-09-30), el primer y último día del mes."""
    mes_txt, anio_txt = str(periodo).strip().lower().split("-")
    mes = _MESES[mes_txt[:3]]
    anio = 2000 + int(anio_txt)
    ini = pd.Timestamp(year=anio, month=mes, day=1)
    fin = ini + pd.offsets.MonthEnd(0)
    return ini, fin


def trimestre_con_anio(trim, periodo):
    """'Q3' -> 'Q3 2026' (el año sale de las fechas del período)."""
    anio = periodo_a_fechas(periodo)[0].year
    t = str(trim).strip()
    return t if str(anio) in t else f"{t} {anio}"


def _fila(ws, contiene, max_row=20):
    for r in range(1, max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and contiene.lower() in str(v).lower():
            return r
    return None


def leer_presupuesto(path_o_bytes):
    """Carga y procesa la hoja de presupuesto mensual y trimestral."""
    if isinstance(path_o_bytes, (str, bytes)):
        wb = openpyxl.load_workbook(path_o_bytes, data_only=True)
    elif hasattr(path_o_bytes, "read"):
        wb = openpyxl.load_workbook(path_o_bytes, data_only=True)
    else:
        wb = openpyxl.load_workbook(str(path_o_bytes), data_only=True)

    try:
        ws = wb.worksheets[0]

        # Trimestre por columna: la fila 1 tiene celdas combinadas
        trimestre = {}
        for rng in ws.merged_cells.ranges:
            if rng.min_row == 1:
                val = ws.cell(row=1, column=rng.min_col).value
                if val:
                    for c in range(rng.min_col, rng.max_col + 1):
                        trimestre[c] = str(val).strip()

        fila_periodo = _fila(ws, "periodo")
        fila_ingresos = _fila(ws, "ingresos")
        if fila_periodo is None or fila_ingresos is None:
            raise ValueError("No encontré las filas 'Periodo' o 'Ingresos' en el archivo.")

        filas = []
        for c in range(2, ws.max_column + 1):
            periodo = ws.cell(row=fila_periodo, column=c).value
            if periodo is None:
                break  # Fin de la tabla principal
            monto = ws.cell(row=fila_ingresos, column=c).value
            filas.append({
                "periodo": str(periodo).strip(),
                "trimestre": trimestre.get(c, "-"),
                "presupuesto_usd": float(monto or 0.0),
            })

        df = pd.DataFrame(filas)

        # El archivo ARDP también trae metas oficiales por año en las columnas
        # posteriores a los 24 meses: "Año 1", "Año 2", etc. Esas metas NO se
        # recalculan sumando los meses, porque el archivo puede usar una meta
        # anual redondeada/determinada por negocio (Año 1 = US$ 1.500.000).
        metas_anuales = {}
        for c in range(2, ws.max_column + 1):
            encabezado = ws.cell(row=2, column=c).value
            if isinstance(encabezado, str) and encabezado.strip().lower().startswith("año "):
                try:
                    numero = int(encabezado.strip().split()[-1])
                    valor = ws.cell(row=fila_ingresos, column=c).value
                    if valor is not None:
                        metas_anuales[numero] = float(valor)
                except (ValueError, TypeError):
                    continue
        df.attrs["metas_anuales"] = metas_anuales
        return df
    finally:
        wb.close()


@st.cache_data(show_spinner=False)
def cargar_presupuesto_cache(ruta_str, mtime=None):
    """Lectura en caché que se invalida automáticamente si el archivo cambia en disco."""
    with open(ruta_str, "rb") as f:
        df = leer_presupuesto(io.BytesIO(f.read()))
    df["trimestre"] = [trimestre_con_anio(t, p) for t, p in zip(df["trimestre"], df["periodo"])]
    return df
