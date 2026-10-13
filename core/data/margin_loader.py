"""Lee el margen real por SKU desde la hoja MATRIZ (TENDIDO_MEZCLA.xlsx o costos.xlsx).
Esa hoja ya trae, por producto (columna 'CODIGO 21' = SKU de Falabella):
costo de fábrica, costo con el factor real (CP COP x9), precio de lista,
comisión, IVA y la rentabilidad (RENT) ya calculada por el jefe.

Fórmulas respetadas estrictamente:
RENT del Excel = 1 − costo_con_factor / (PV / 1,19)
"""
import io
import re
import unicodedata
import openpyxl
import pandas as pd
import streamlit as st


def _norm(x):
    """'Comisión ' -> 'COMISION'; None -> ''."""
    if x is None:
        return ""
    s = unicodedata.normalize("NFKD", str(x))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().upper()


_COLS = {
    "REF": "ref", "CP $ USD": "costo_fabrica_usd", "CP $ COP": "costo_fabrica_cop",
    "COSTO CON FACTOR": "costo_real_cop", "PV": "precio_lista_cop", "RENT": "rent_lista",
    "COMISION": "comision_cop", "IVA": "iva_cop", "RENT FINAL": "rent_final_cop",
    "CODIGO 21": "sku", "NAMING": "producto",
}
_REQUERIDAS = ["ref", "sku", "costo_real_cop", "rent_lista"]
_MAX_FILAS_BUSCANDO_HEADER = 40


def _buscar_hoja(wb, nombre="MATRIZ"):
    for n in wb.sheetnames:
        if _norm(n) == nombre:
            return wb[n]
    raise ValueError(f"No encontré la hoja '{nombre}'. Hojas del archivo: {wb.sheetnames}")


def _mapear_header(header):
    col = {}
    for i, h in enumerate(header):
        n = _norm(h)
        if n in _COLS and _COLS[n] not in col:
            col[_COLS[n]] = i
    if "costo_real_cop" not in col:
        for i, h in enumerate(header):
            if "FACTOR" in _norm(h):
                col["costo_real_cop"] = i
                break
    return col


def _sku_texto(v):
    try:
        f = float(v)
        if f.is_integer():
            return str(int(f))
    except (TypeError, ValueError):
        pass
    return str(v).strip()


def leer_margenes(path_o_bytes):
    """Lee y limpia la matriz de costos y rentabilidad de lista."""
    wb = openpyxl.load_workbook(path_o_bytes, data_only=True, read_only=True)
    try:
        ws = _buscar_hoja(wb)
        col = None
        ultimo_header = []
        filas = []
        vacias_seguidas = 0

        for n_fila, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if col is None:
                if n_fila > _MAX_FILAS_BUSCANDO_HEADER:
                    raise ValueError(
                        "No encontré la fila de encabezados (con 'REF' y 'CODIGO 21') en las "
                        f"primeras {_MAX_FILAS_BUSCANDO_HEADER} filas de la hoja MATRIZ."
                    )
                posible = _mapear_header(row)
                if "ref" in posible and "sku" in posible:
                    col = posible
                    ultimo_header = [_norm(h) for h in row if _norm(h)]
                    faltan = [c for c in _REQUERIDAS if c not in col]
                    if faltan:
                        raise ValueError(
                            f"En la hoja MATRIZ (fila {n_fila}) faltan las columnas {faltan}. "
                            f"Encabezados encontrados: {ultimo_header}"
                        )
                continue

            ref = row[col["ref"]] if col["ref"] < len(row) else None
            sku = row[col["sku"]] if col["sku"] < len(row) else None
            if not ref or sku is None:
                vacias_seguidas += 1
                if vacias_seguidas > 50:
                    break
                continue
            vacias_seguidas = 0

            fila = {"ref": str(ref).strip(), "sku": _sku_texto(sku)}
            for nombre, idx in col.items():
                if nombre in ("ref", "sku"):
                    continue
                fila[nombre] = row[idx] if idx < len(row) else None
            filas.append(fila)
    finally:
        wb.close()

    if col is None:
        raise ValueError("La hoja MATRIZ está vacía o no tiene encabezados reconocibles.")
    if not filas:
        raise ValueError("Encontré los encabezados pero ninguna fila con REF y CODIGO 21 llenos.")

    df = pd.DataFrame(filas)
    for c in _COLS.values():
        if c not in df.columns:
            df[c] = pd.NA
    for c in ("costo_real_cop", "rent_lista", "precio_lista_cop"):
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    df["margen_lista_%"] = df["rent_lista"] * 100
    return df.drop_duplicates(subset="sku")


@st.cache_data(show_spinner=False)
def cargar_margenes_cache(ruta_str, mtime=None):
    with open(ruta_str, "rb") as f:
        return leer_margenes(io.BytesIO(f.read()))
